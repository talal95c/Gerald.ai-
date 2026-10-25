from __future__ import annotations

from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor
import logging
from pathlib import Path
import shutil

import numpy as np

from gerald.library.store import Document, Store

logger = logging.getLogger(__name__)


def extract_text(path: str | Path) -> list[tuple[int | None, str]]:
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(source))
        return [
            (page_number, page.extract_text() or "")
            for page_number, page in enumerate(reader.pages, start=1)
        ]
    if suffix == ".docx":
        from docx import Document as DocxDocument

        document = DocxDocument(str(source))
        text = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        text.extend(
            "\t".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return [(None, "\n".join(text))]
    if suffix in {".txt", ".md"}:
        return [(None, source.read_text(encoding="utf-8", errors="replace"))]
    if suffix in {".png", ".jpg", ".jpeg"}:
        from PIL import Image

        from gerald.perception.elements import ocr_image_text

        with Image.open(source) as image:
            text = ocr_image_text(image.convert("RGB"))
        return [(None, text)]
    raise ValueError("Format non pris en charge")


def chunk_pages(
    pages: list[tuple[int | None, str]],
    size: int = 1200,
    overlap: int = 200,
) -> list[tuple[int | None, str]]:
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("size must be positive and overlap must be smaller than size")
    chunks: list[tuple[int | None, str]] = []
    for page, page_text in pages:
        text = page_text.strip()
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                lower_bound = start + max(1, size * 3 // 5)
                boundary = max(
                    text.rfind("\n\n", lower_bound, end),
                    text.rfind("\n", lower_bound, end),
                    text.rfind(" ", lower_bound, end),
                )
                if boundary > start:
                    end = boundary
            chunk = text[start:end].strip()
            if chunk:
                chunks.append((page, chunk))
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
            while start < len(text) and text[start].isspace():
                start += 1
    return chunks


class Indexer:
    def __init__(self, store: Store, embedder, workers: int = 2) -> None:
        self.store = store
        self.embedder = embedder
        self._executor = ThreadPoolExecutor(
            max_workers=workers,
            thread_name_prefix="gerald-indexer",
        )

    def add_document(
        self,
        course_id: int,
        path: str | Path,
        on_done: Callable[[Document], None] | None = None,
    ) -> Document:
        source = Path(path)
        document = self.store._create_document(course_id, source.name)
        self._executor.submit(self._index_document, document, source, on_done)
        return document

    def _index_document(
        self,
        document: Document,
        source: Path,
        on_done: Callable[[Document], None] | None,
    ) -> None:
        updated = None
        try:
            destination = Path(document.stored_path)
            if source.resolve() != destination.resolve():
                shutil.copy2(source, destination)
            chunks = chunk_pages(extract_text(destination))
            stored_chunks: list[tuple[int | None, str, bytes | None]] = []
            embedding_available = True
            warned = False
            for start in range(0, len(chunks), 64):
                batch = chunks[start : start + 64]
                vectors = None
                if embedding_available:
                    try:
                        result = np.asarray(
                            self.embedder.embed([text for _page, text in batch]),
                            dtype=np.float32,
                        )
                        if result.ndim != 2 or len(result) != len(batch):
                            raise ValueError("Embedding response did not match the document batch")
                        vectors = result
                    except Exception:
                        embedding_available = False
                        if not warned:
                            logger.warning(
                                "Embeddings unavailable; indexing document without remaining vectors"
                            )
                            warned = True
                for index, (page, text) in enumerate(batch):
                    embedding = (
                        vectors[index].astype(np.float32, copy=False).tobytes()
                        if vectors is not None
                        else None
                    )
                    stored_chunks.append((page, text, embedding))
            self.store._replace_chunks(document.id, stored_chunks)
            updated = self.store._finish_document(
                document.id,
                status="ready",
                error=None,
                n_chunks=len(stored_chunks),
            )
        except Exception as error:
            updated = self.store._finish_document(
                document.id,
                status="error",
                error=str(error),
                n_chunks=0,
            )
        if on_done is not None and updated is not None:
            try:
                on_done(updated)
            except Exception:
                logger.warning("Document indexing callback failed", exc_info=True)

    def close(self) -> None:
        self._executor.shutdown(wait=False, cancel_futures=True)
