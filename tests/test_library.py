import json
import threading
from pathlib import Path

import numpy as np
from PIL import Image
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

import gerald.library as library
from gerald.library import embeddings, qa
from gerald.library.ingest import Indexer, chunk_pages, extract_text
from gerald.library.search import Hit, search
from gerald.library.store import Store


def _write_text_pdf(path, pages):
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = writer._add_object(
            DictionaryObject(
                {
                    NameObject("/Type"): NameObject("/Font"),
                    NameObject("/Subtype"): NameObject("/Type1"),
                    NameObject("/BaseFont"): NameObject("/Helvetica"),
                }
            )
        )
        resources = writer._add_object(
            DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
                }
            )
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


def test_store_crud_cascade_and_imported_file_cleanup(tmp_path):
    store = Store(tmp_path / "library")
    assert store.get_profile() is None
    profile = store.save_profile("Camille", "Lycée", ["Maths", "Physique"])
    assert profile.name == "Camille"
    assert store.get_profile().subjects == ["Maths", "Physique"]

    course = store.create_course("Physique", "Physique")
    course = store.update_course(course.id, title="Mécanique", notes="Forces")
    assert course.title == "Mécanique"
    assert course.notes == "Forces"

    source = tmp_path / "cours.txt"
    source.write_text("La force accélère un objet.", encoding="utf-8")

    class NoEmbeddings:
        def embed(self, _texts):
            raise RuntimeError("offline")

    completed = threading.Event()
    callback_rows = []
    callback_threads = []
    indexer = Indexer(store, NoEmbeddings())

    def on_done(row):
        callback_rows.append(row)
        callback_threads.append(threading.current_thread().name)
        completed.set()

    pending = indexer.add_document(
        course.id,
        source,
        on_done,
    )
    assert pending.status == "indexing"
    assert completed.wait(5)
    indexer.close()
    document = store.list_documents(course.id)[0]
    assert document.status == "ready"
    assert document.n_chunks == 1
    assert Path(document.stored_path).read_text(encoding="utf-8") == source.read_text(encoding="utf-8")
    assert callback_rows[0].status == "ready"
    assert callback_threads[0].startswith("gerald-indexer")

    conversation = store.create_conversation(course.id, "chat", "Les forces")
    message = store.add_message(conversation.id, "assistant", "Une force agit.", [{"page": 1}])
    assert store.get_messages(conversation.id) == [message]
    recent = store.recent_conversations(source="chat", course_id=course.id)
    assert [item.id for item in recent] == [conversation.id]

    store.delete_course(course.id)
    assert store.list_courses() == []
    assert store.list_documents(course.id) == []
    assert not Path(document.stored_path).exists()
    assert store.get_conversation(conversation.id).course_id is None
    assert store.get_messages(conversation.id)[0].sources == [{"page": 1}]
    store.close()


def test_store_deduplicates_document_filenames_and_delete_document_removes_file(tmp_path):
    store = Store(tmp_path / "library")
    course = store.create_course("Histoire", "Histoire")
    first = store._create_document(course.id, "notes.txt")
    second = store._create_document(course.id, "notes.txt")
    assert first.filename == "notes.txt"
    assert second.filename == "notes (2).txt"
    Path(first.stored_path).write_text("contenu", encoding="utf-8")

    store.delete_document(first.id)

    assert not Path(first.stored_path).exists()
    assert [item.id for item in store.list_documents(course.id)] == [second.id]
    store.close()


def test_extract_text_from_txt_markdown_docx_and_pdf(tmp_path, monkeypatch):
    text = tmp_path / "cours.txt"
    text.write_bytes(b"Bonjour \xff")
    markdown = tmp_path / "notes.md"
    markdown.write_text("# Leçon\n\nUn concept.", encoding="utf-8")
    assert extract_text(text) == [(None, "Bonjour \ufffd")]
    assert extract_text(markdown) == [(None, "# Leçon\n\nUn concept.")]

    docx_path = tmp_path / "cours.docx"
    document = DocxDocument()
    document.add_paragraph("La matière est composée d'atomes.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Proton"
    table.cell(0, 1).text = "Charge positive"
    document.save(docx_path)
    docx_text = extract_text(docx_path)[0][1]
    assert "La matière est composée d'atomes." in docx_text
    assert "Proton\tCharge positive" in docx_text

    pdf_path = tmp_path / "cours.pdf"
    _write_text_pdf(pdf_path, ["Premiere page", "Deuxieme page"])
    pdf_text = extract_text(pdf_path)
    assert pdf_text == [(1, "Premiere page"), (2, "Deuxieme page")]

    image_path = tmp_path / "cours.png"
    Image.new("RGB", (40, 30), "white").save(image_path)
    monkeypatch.setattr("gerald.perception.elements.ocr_image_text", lambda _image: "Texte OCR")
    assert extract_text(image_path) == [(None, "Texte OCR")]


def test_extract_text_rejects_unsupported_format(tmp_path):
    path = tmp_path / "cours.rtf"
    path.write_text("unsupported", encoding="utf-8")

    with pytest.raises(ValueError, match="Format non pris en charge"):
        extract_text(path)


def test_chunking_keeps_page_numbers_overlap_and_paragraph_boundaries():
    chunks = chunk_pages(
        [(3, "Premier paragraphe.\n\n" + ("mot " * 100))],
        size=120,
        overlap=20,
    )

    assert chunks
    assert all(page == 3 for page, _text in chunks)
    assert any("Premier paragraphe." in text for _page, text in chunks)
    assert all(len(text) <= 120 for _page, text in chunks)


def test_indexer_records_unsupported_format_as_error_and_calls_back(tmp_path):
    store = Store(tmp_path / "library")
    course = store.create_course("Sciences", "Sciences")
    source = tmp_path / "cours.rtf"
    source.write_text("incompatible", encoding="utf-8")
    completed = threading.Event()
    result = []
    indexer = Indexer(store, object())
    pending = indexer.add_document(course.id, source, lambda row: (result.append(row), completed.set()))

    assert pending.status == "indexing"
    assert completed.wait(5)
    assert result[0].status == "error"
    assert result[0].error == "Format non pris en charge"
    indexer.close()
    store.close()


def test_hybrid_search_ranks_embedded_match_and_keyword_only_fallback(tmp_path):
    store = Store(tmp_path / "library")
    course = store.create_course("Biologie", "Biologie")
    files = []
    for filename, content in (
        ("energie.txt", "L'énergie se conserve dans un système."),
        ("cellule.txt", "La cellule végétale réalise la photosynthèse."),
    ):
        path = tmp_path / filename
        path.write_text(content, encoding="utf-8")
        files.append(path)

    class TopicEmbedder:
        def embed(self, texts):
            return np.asarray(
                [[0.0, 1.0] if "cellule" in text or "photosynthèse" in text else [1.0, 0.0] for text in texts],
                dtype=np.float32,
            )

    indexer = Indexer(store, TopicEmbedder(), workers=2)
    completed = [threading.Event(), threading.Event()]
    for index, path in enumerate(files):
        indexer.add_document(course.id, path, lambda _row, event=completed[index]: event.set())
    assert all(event.wait(5) for event in completed)
    hits = search(store, "photosynthèse", TopicEmbedder(), course.id)
    assert hits[0].document == "cellule.txt"

    class OfflineEmbedder:
        def embed(self, _texts):
            raise RuntimeError("offline")

    keyword_hits = search(store, "énergie système", OfflineEmbedder(), course.id)
    assert keyword_hits[0].document == "energie.txt"

    offline_source = tmp_path / "terrain.txt"
    offline_source.write_text("Le calcaire est une roche.", encoding="utf-8")
    offline_done = threading.Event()
    offline_indexer = Indexer(store, OfflineEmbedder())
    offline_indexer.add_document(course.id, offline_source, lambda _row: offline_done.set())
    assert offline_done.wait(5)
    no_embedding_hits = search(store, "calcaire roche", OfflineEmbedder(), course.id)
    assert no_embedding_hits[0].document == "terrain.txt"
    indexer.close()
    offline_indexer.close()
    store.close()


def test_keyword_only_search_drops_chunks_without_overlap(tmp_path):
    store = Store(tmp_path / "library")
    course = store.create_course("Géographie", "Géographie")
    sources = {
        "loire.txt": "Le fleuve Loire traverse Orléans.",
        "mitochondries.txt": "Les mitochondries fournissent l'énergie cellulaire.",
        "volcans.txt": "Les volcans émettent de la lave.",
    }

    class OfflineEmbedder:
        def embed(self, _texts):
            raise RuntimeError("offline")

    embedder = OfflineEmbedder()
    indexer = Indexer(store, embedder, workers=3)
    completed = []
    for filename, text in sources.items():
        path = tmp_path / filename
        path.write_text(text, encoding="utf-8")
        event = threading.Event()
        completed.append(event)
        indexer.add_document(course.id, path, lambda _row, event=event: event.set())
    assert all(event.wait(5) for event in completed)

    hits = search(store, "Loire", embedder, course.id)
    assert [hit.document for hit in hits] == ["loire.txt"]
    assert search(store, "Atlantique", embedder, course.id) == []

    indexer.close()
    store.close()


def test_openai_embedder_sends_versioned_user_agent_and_returns_float32(monkeypatch):
    captured = {}

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def read(self):
            return b'{"data":[{"index":0,"embedding":[0.25,0.75]}]}'

    def fake_urlopen(request, timeout):
        captured["request"] = request
        captured["timeout"] = timeout
        return Response()

    monkeypatch.setattr(embeddings.urllib.request, "urlopen", fake_urlopen)
    vector = library.OpenAIEmbedder("test-key").embed(["un passage"])
    request = captured["request"]
    headers = {key.casefold(): value for key, value in request.header_items()}

    assert request.full_url == "https://api.openai.com/v1/embeddings"
    assert headers["user-agent"] == "gerald/0.1.0"
    assert json.loads(request.data) == {
        "model": "text-embedding-3-small",
        "input": ["un passage"],
    }
    assert captured["timeout"] == 30
    assert vector.dtype == np.float32
    assert vector.tolist() == [[0.25, 0.75]]


def test_qa_sends_profile_sources_and_last_ten_history_and_maps_citations(monkeypatch):
    captured = {}

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def read(self):
            return b'{"output":[{"type":"message","content":[{"type":"output_text","text":"Le cycle [1] aide la plante."}]}]}'

    def fake_urlopen(request, timeout):
        captured["request"] = request
        captured["timeout"] = timeout
        return Response()

    monkeypatch.setattr(qa.urllib.request, "urlopen", fake_urlopen)
    source = Hit("Biologie", "cours.pdf", 3, "La lumière aide la plante.", 0.9)
    history = [{"role": "user", "text": f"ancienne question {index}"} for index in range(11)]

    result = library.answer(
        "Pourquoi la lumière est-elle utile ?",
        [source],
        history,
        {"name": "Léa", "level": "Lycée", "subjects": ["SVT"]},
        "test-key",
    )

    request = captured["request"]
    headers = {key.casefold(): value for key, value in request.header_items()}
    payload = json.loads(request.data)
    assert request.full_url == "https://api.openai.com/v1/responses"
    assert headers["user-agent"] == "gerald/0.1.0"
    assert captured["timeout"] == 60
    assert payload["model"] == "gpt-5.4-mini"
    assert payload["reasoning"] == {"effort": "low"}
    assert "Léa" in payload["instructions"]
    assert "Lycée" in payload["instructions"]
    assert "[1] cours.pdf, page 3" in payload["input"][-1]["content"]
    assert len(payload["input"]) == 11
    assert payload["input"][0]["content"] == "ancienne question 1"
    assert result.text == "Le cycle [1] aide la plante."
    assert result.sources == [source]


def test_qa_only_sends_reasoning_for_gpt5_models(monkeypatch):
    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def read(self):
            return b'{"output":[]}'

    captured = {}

    def fake_urlopen(request, **_kwargs):
        captured["payload"] = json.loads(request.data)
        return Response()

    monkeypatch.setattr(qa.urllib.request, "urlopen", fake_urlopen)
    library.answer("Question", [], [], None, "test-key", model="gpt-4.1-mini")

    assert "reasoning" not in captured["payload"]


def test_ask_persists_chat_messages_and_cited_sources(monkeypatch, tmp_path):
    store = Store(tmp_path / "library")
    course = store.create_course("Physique", "Physique")
    source = Hit("Physique", "forces.txt", None, "Une force accélère un objet.", 0.8)

    def fake_answer(question, hits, history, profile, api_key, model):
        assert question == "Pourquoi l'objet accélère ?"
        assert hits == []
        assert history == []
        assert profile is None
        assert api_key == "test-key"
        assert model is None
        return library.Answer("Parce qu'une force agit. [1]", [source])

    monkeypatch.setattr(library, "answer", fake_answer)
    conversation_id, result = library.ask(
        store,
        None,
        "Pourquoi l'objet accélère ?",
        None,
        course.id,
        "test-key",
    )

    conversation = store.get_conversation(conversation_id)
    messages = store.get_messages(conversation_id)
    assert conversation.source == "chat"
    assert conversation.course_id == course.id
    assert conversation.title == "Pourquoi l'objet accélère ?"
    assert [message.role for message in messages] == ["user", "assistant"]
    assert messages[1].sources == [
        {
            "course_title": "Physique",
            "document": "forces.txt",
            "page": None,
            "text": "Une force accélère un objet.",
            "score": 0.8,
        }
    ]
    assert result.sources == [source]
    store.close()
